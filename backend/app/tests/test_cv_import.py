"""CV import: which fields a CV may fill, the candidate check (verbatim
quote, accepted value), the upload rules (consent, file type, readable
text, rate limit, extraction failure), and how reviewed answers are stored
(CV for an unchanged candidate, QA otherwise; conflicts for fields answered
differently) — on in-memory stand-ins for the profile store, the cache and
the extraction call."""

from __future__ import annotations

import io

import pytest
from docx import Document

from app.core.errors import ForbiddenError, NotFoundError, RateLimitError, ServiceUnavailableError, ValidationError
from app.domains.cv_import import cv_text
from app.domains.cv_import import service as cv_service
from app.domains.profile import service as profile_service
from app.tests.profile_fixtures import DEFINITIONS, ROLE_TAGS, entry

# The fixture dictionary with extraction types as in the live one.
CV_DEFINITIONS = {
    key: {**d, "extraction_type": "literal" if "CV" in d["sources"] else None}
    for key, d in DEFINITIONS.items()
}

CV_TEXT = (
    "Jane Tan\nBachelor of Engineering (Computer Engineering), National University of Singapore, 2021\n"
    "Certifications: CFA Level II candidate; FRM Part I passed\n"
    "Skills: Python, SQL, Excel\n" + "Experience at a bank as a data analyst working on risk dashboards. " * 4
)


class _Store:
    """profile repository stand-in: the profile, conflicts, CV documents."""

    def __init__(self) -> None:
        self.profile = {"lifecycle_stage": entry("prospect")}
        self.conflicts: list[dict] = []
        self.closed: list[tuple] = []
        self.documents: dict[str, dict] = {}
        self.writes: list[dict] = []

    def get_profile(self, user_id):
        return dict(self.profile)

    def list_open_conflicts(self, user_id):
        return [c for c in self.conflicts if c["status"] == "open"]

    def record_conflict(self, user_id, key, *, kind, current_value, current_source, candidate_value, candidate_source):
        self.conflicts.append({"id": len(self.conflicts) + 1, "field_key": key, "kind": kind, "status": "open",
                               "current_value": current_value, "candidate_value": candidate_value})
        return len(self.conflicts)

    def close_conflict(self, conflict_id, status):
        for c in self.conflicts:
            if c["id"] == conflict_id:
                c["status"] = status
        self.closed.append((conflict_id, status))
        return True

    def create_cv_document(self, user_id, *, storage_bucket, storage_path, retain_original=False):
        doc_id = f"doc{len(self.documents) + 1}"
        self.documents[doc_id] = {"id": doc_id, "user_id": user_id, "parse_status": "pending", "candidates": {},
                                  "error": None, "uploaded_at": None, "parsed_at": None,
                                  "storage_path": storage_path}
        return self.documents[doc_id]

    def save_cv_result(self, document_id, *, candidates=None, error=None):
        doc = self.documents[document_id]
        doc.update(parse_status="failed" if error else "parsed", candidates=candidates or {}, error=error)
        return True

    def get_cv_document(self, user_id, document_id):
        doc = self.documents.get(document_id)
        return doc if doc and doc["user_id"] == user_id else None

    def get_latest_cv_document(self, user_id):
        docs = [d for d in self.documents.values() if d["user_id"] == user_id]
        return docs[-1] if docs else None

    def set_fields(self, user_id, changes, *, source, actor, from_cv=None):
        self.writes.append({"changes": changes, "from_cv": from_cv or {}})
        for k, v in changes.items():
            self.profile[k] = entry(v, "CV" if k in (from_cv or {}) else source)


class _Cache:
    def __init__(self) -> None:
        self.data: dict[str, str] = {}

    def get(self, key):
        return self.data.get(key)

    def set(self, key, value, ttl_seconds):
        self.data[key] = value


@pytest.fixture()
def store(monkeypatch):
    fake = _Store()
    r = profile_service.repository
    monkeypatch.setattr(r, "load_field_definitions", lambda: CV_DEFINITIONS)
    monkeypatch.setattr(profile_service, "_role_tags", lambda: ROLE_TAGS)
    for name in ("get_profile", "list_open_conflicts", "record_conflict", "close_conflict", "create_cv_document",
                 "save_cv_result", "get_cv_document", "get_latest_cv_document", "set_fields"):
        monkeypatch.setattr(r, name, getattr(fake, name))
    monkeypatch.setattr(r, "has_consent", lambda user_id, key: True)
    monkeypatch.setattr(r, "list_cv_quotes", lambda user_id: {})
    monkeypatch.setattr(cv_service, "cache", _Cache())
    return fake


def _docx(text: str) -> bytes:
    document = Document()
    for line in text.split("\n"):
        document.add_paragraph(line)
    table = document.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "Languages: English, Mandarin"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- targets and the candidate check -----------------------------------------

def test_targets_are_literal_cv_fields_for_the_stage_without_self_ratings(store):
    keys = {t["field_key"] for t in profile_service.cv_targets("u1")}
    assert keys == {"languages", "bachelor_grad_year", "degree_discipline_category", "admission_tests",
                    "certifications", "current_position"}
    assert "code_level" not in keys  # a self-rating, even though CV is a source


def test_targets_follow_the_lifecycle_stage(store):
    store.profile["lifecycle_stage"] = entry("current")  # fixture fields apply to prospects only
    assert profile_service.cv_targets("u1") == []


def test_candidates_need_a_verbatim_quote_and_an_accepted_value(store):
    targets = profile_service.cv_targets("u1")
    raw = {
        "degree_discipline_category": {"value": "engineering",
                                       "quote": "bachelor of engineering   (Computer Engineering)"},
        "certifications": {"value": ["frm", "cfa"], "quote": "CFA Level II candidate; FRM Part I passed"},
        "bachelor_grad_year": {"value": 2021, "quote": "graduated with first class honours"},  # not in the CV
        "languages": {"value": ["en"], "quote": ""},
        "code_level": {"value": 3, "quote": "Python, SQL"},  # not a target
        "admission_tests": {"value": [{"key": "SAT", "score": 1500}], "quote": "Jane Tan"},  # bad value
    }
    candidates = cv_service.check_candidates(raw, targets, CV_TEXT)
    assert set(candidates) == {"degree_discipline_category", "certifications"}
    assert candidates["degree_discipline_category"]["quote"] == "bachelor of engineering   (Computer Engineering)"


def test_docx_text_includes_tables():
    text = cv_text.extract_text(_docx("Jane Tan\nPython"), "cv.DOCX")
    assert "Jane Tan" in text and "Languages: English, Mandarin" in text


def test_only_pdf_and_docx_are_accepted():
    with pytest.raises(ValidationError):
        cv_text.extract_text(b"hello", "cv.txt")
    with pytest.raises(ValidationError):
        cv_text.extract_text(b"not a pdf", "cv.pdf")


# --- upload ----------------------------------------------------------------------

def _extract_returns(monkeypatch, raw):
    calls = []

    def fake(text, targets):
        calls.append(targets)
        if isinstance(raw, Exception):
            raise raw
        return raw
    monkeypatch.setattr(cv_service, "extract_fields", fake)
    return calls


def test_upload_needs_cv_consent(store, monkeypatch):
    monkeypatch.setattr(profile_service.repository, "has_consent", lambda user_id, key: False)
    with pytest.raises(ForbiddenError):
        cv_service.import_cv("u1", "cv.docx", _docx(CV_TEXT))


def test_unreadable_files_are_rejected_before_extraction(store, monkeypatch):
    calls = _extract_returns(monkeypatch, {})
    with pytest.raises(ValidationError):
        cv_service.import_cv("u1", "cv.docx", _docx("Jane Tan"))  # too little text, e.g. a scan
    with pytest.raises(ValidationError):
        cv_service.import_cv("u1", "cv.docx", b"x" * (cv_service.MAX_FILE_BYTES + 1))
    assert calls == [] and store.documents == {}


def test_import_stores_checked_candidates_and_flags_conflicts(store, monkeypatch):
    store.profile["certifications"] = entry(["other"])  # answered differently before
    _extract_returns(monkeypatch, {
        "degree_discipline_category": {"value": "engineering", "quote": "Bachelor of Engineering"},
        "certifications": {"value": ["cfa"], "quote": "CFA Level II candidate"},
        "bachelor_grad_year": {"value": 2019, "quote": "made up"},
    })
    result = cv_service.import_cv("u1", "cv.docx", _docx(CV_TEXT))
    statuses = {f["field_key"]: f["status"] for f in result["fields"]}
    assert statuses == {"degree_discipline_category": "new", "certifications": "different"}
    assert result["parse_status"] == "parsed"
    assert store.documents["doc1"]["storage_path"] is None  # the file isn't kept
    assert [(c["field_key"], c["kind"]) for c in store.conflicts] == [("certifications", "self_vs_cv")]

    cv_service.import_cv("u1", "cv.docx", _docx(CV_TEXT))  # same candidate again: no second conflict
    assert len(store.conflicts) == 1


def test_failed_extraction_is_recorded_and_reported(store, monkeypatch):
    _extract_returns(monkeypatch, ValueError("not JSON"))
    with pytest.raises(ServiceUnavailableError):
        cv_service.import_cv("u1", "cv.docx", _docx(CV_TEXT))
    assert store.documents["doc1"]["parse_status"] == "failed"


def test_uploads_are_rate_limited(store, monkeypatch):
    _extract_returns(monkeypatch, {})
    for _ in range(cv_service.UPLOADS_PER_HOUR):
        cv_service.import_cv("u1", "cv.docx", _docx(CV_TEXT))
    with pytest.raises(RateLimitError):
        cv_service.import_cv("u1", "cv.docx", _docx(CV_TEXT))


# --- saving reviewed answers ---------------------------------------------------------

def test_unchanged_candidates_are_stored_as_cv_and_settle_conflicts(store, monkeypatch):
    store.profile["certifications"] = entry(["other"])
    _extract_returns(monkeypatch, {
        "degree_discipline_category": {"value": "engineering", "quote": "Bachelor of Engineering"},
        "certifications": {"value": ["cfa", "frm"], "quote": "CFA Level II candidate; FRM Part I passed"},
    })
    doc_id = cv_service.import_cv("u1", "cv.docx", _docx(CV_TEXT))["document_id"]

    profile_service.update_fields(
        "u1", {"degree_discipline_category": "cs_se", "certifications": ["frm", "cfa"]}, cv_document_id=doc_id,
    )
    write = store.writes[-1]
    assert set(write["from_cv"]) == {"certifications"}  # the edited discipline stays QA
    assert write["from_cv"]["certifications"]["quote"] == "CFA Level II candidate; FRM Part I passed"
    assert store.profile["certifications"]["source"] == "CV"
    assert store.profile["degree_discipline_category"]["source"] == "QA"
    assert store.closed == [(1, "resolved")]


def test_saving_against_someone_elses_import_is_not_found(store):
    with pytest.raises(NotFoundError):
        profile_service.update_fields("u1", {"degree_discipline_category": "cs_se"}, cv_document_id="nope")


def test_settings_take_two_letter_language_codes():
    from pydantic import ValidationError as PydanticValidationError

    from app.domains.auth.schemas import SettingsUpdate

    assert SettingsUpdate(language_pref=["EN", "zh", "en"]).language_pref == ["en", "zh"]
    with pytest.raises(PydanticValidationError):
        SettingsUpdate(language_pref=["english"])
    with pytest.raises(PydanticValidationError):
        SettingsUpdate(answer_depth="verbose")

