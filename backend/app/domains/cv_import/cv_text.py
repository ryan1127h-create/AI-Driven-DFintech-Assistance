"""CV file -> plain text, for PDF and DOCX. No OCR: a scanned PDF yields
(almost) no text, which the service reports as unreadable."""

from __future__ import annotations

import io

from app.core.errors import ValidationError

SUPPORTED_EXTENSIONS = ("pdf", "docx")


def extract_text(data: bytes, filename: str) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValidationError("Upload your CV as a PDF or DOCX file.")
    try:
        return _pdf_text(data) if ext == "pdf" else _docx_text(data)
    except Exception as exc:
        raise ValidationError("This file couldn't be opened. Check it isn't damaged or password-protected.") from exc


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_text(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs]
    # CV templates often lay sections out in tables.
    for table in document.tables:
        for row in table.rows:
            lines.extend(cell.text for cell in row.cells)
    return "\n".join(lines)
